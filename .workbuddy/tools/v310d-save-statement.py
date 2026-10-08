# -*- coding: utf-8 -*-
"""v310d 落盘简爱对账单附件（base64 -> xlsx）+ 结构 dump"""
import base64
import os
import re
import zipfile
import warnings
import openpyxl

warnings.filterwarnings("ignore")

OUT = "/Users/zhangjunfeng/Documents/流水对账/厂商对账单样本"
os.makedirs(OUT, exist_ok=True)

B64 = open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "v310d-b64.txt"), "r").read().strip()
raw = base64.b64decode(B64)
p = os.path.join(OUT, "简爱-华中大区-湖北福宝商贸有限公司2026年8月销售对账单.xlsx")
open(p, "wb").write(raw)
print("已落盘:", p, "%.1f KB" % (len(raw) / 1024))

z = zipfile.ZipFile(p)
ns = z.namelist()
print("zip 部件 %d 个；透视表部件:" % len(ns), [n for n in ns if "pivot" in n.lower()])
x = z.read("xl/workbook.xml").decode("utf-8", "ignore")
for m in re.finditer(r'<sheet[^>]*name="([^"]*)"', x):
    print("  sheet:", m.group(1))
z.close()

wb = openpyxl.load_workbook(p, data_only=True)
for ws in wb.worksheets:
    print()
    print("=" * 78)
    print("sheet「%s」 max_row=%s max_col=%s" % (ws.title, ws.max_row, ws.max_column))
    print("=" * 78)
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(18, ws.max_row or 0),
                                         max_col=min(ws.max_column or 0, 20))):
        vals = ["" if c.value is None else str(c.value).replace("\n", "\\n").strip()
                for c in row]
        while vals and vals[-1] == "":
            vals.pop()
        print(" R%-2d | %s" % (i + 1, " | ".join(vals)))
wb.close()
