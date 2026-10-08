# -*- coding: utf-8 -*-
"""v310 侦察：舟谱导出三类明细表结构 dump
只读，不改任何原文件。
"""
import os
import sys
import zipfile
import openpyxl

D = "/Users/zhangjunfeng/Documents/流水对账/舟谱导出的单据"
FILES = [
    "20260801-20260831采购单明细.xlsx",
    "2026年8月收入明细表.xlsx",
    "2026年8月费用明细表.xlsx",
]

SEP = "=" * 78

def dump_zip(name):
    p = os.path.join(D, name)
    z = zipfile.ZipFile(p)
    ns = z.namelist()
    print("  [zip 内部件] %d 个" % len(ns))
    piv = [n for n in ns if "pivot" in n.lower()]
    print("  透视表相关:", piv if piv else "无")
    sheets = [n for n in ns if n.startswith("xl/worksheets/sheet")]
    print("  工作表部件:", sheets)
    for n in ns:
        if n == "xl/workbook.xml":
            import re
            x = z.read(n).decode("utf-8", "ignore")
            for m in re.finditer(r'<sheet[^>]*name="([^"]*)"[^>]*r:id="([^"]*)"', x):
                print("    sheet: %-24s %s" % (m.group(1), m.group(2)))
    z.close()

def cell_str(v):
    if v is None:
        return ""
    s = str(v).replace("\n", "\\n").strip()
    return s

def dump_sheet(ws, max_rows=8, max_cols=30):
    try:
        dim = ws.calculate_dimension()
    except Exception:
        dim = "?"
    print("  --- sheet「%s」 dim=%s  max_row=%s max_col=%s ---" % (
        ws.title, dim, ws.max_row, ws.max_column))
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row or 0),
                                         max_col=min(max_cols, ws.max_column or 0))):
        vals = [cell_str(c.value) for c in row]
        # 去掉尾部空
        while vals and vals[-1] == "":
            vals.pop()
        print("   R%-2d | %s" % (i + 1, " | ".join(vals)))

def main():
    for name in FILES:
        p = os.path.join(D, name)
        print(SEP)
        print("文件:", name, " (%.1f KB)" % (os.path.getsize(p) / 1024))
        print(SEP)
        dump_zip(name)
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            print("  打开方式: data_only=True  sheetnames=", wb.sheetnames)
            for ws in wb.worksheets:
                dump_sheet(ws)
            wb.close()
        except Exception as e:
            print("  !! 读取失败:", type(e).__name__, e)
        print()

if __name__ == "__main__":
    main()
