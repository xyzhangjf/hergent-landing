#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v184 真机验证 · 步 1：用真实模版造一份导入文件，跑 /api/import/preview，落盘原始响应。

造文件原则：**只用真实模版自身的列**（含租户动态追加的客户列），
这样被测的就是「用户下载模版→填→上传」的真实链路，而不是我手搓的一个理想表头。
"""
import json
import os
import shutil
import sys
import urllib.request

from openpyxl import load_workbook

BASE = "https://hergent.cn"
TOKEN = "a21d9659fd4346c18cd56cf60a57f7eebfbc2bfe21f24b5db5cca07c60133fd4"
TENANT = "9999"

TPL = "/tmp/tpl_cross_9999.xlsx"
OUT = "/tmp/imp_cross_9999.xlsx"

# (条码, 商品名, 规格, 单位, 到货周期填值, 备注)  —— 到货周期为空串 = 留空
ROWS = [
    ("6923644201203", "蒙牛小鲜语绿莎莎牛乳绿豆奶PET瓶185ml×24瓶", "185ml×24瓶", "瓶", "+4天", "标准写法"),
    ("6934665011023", "蒙牛双拼果蔬风味酸牛乳圆周杯90g×8杯×12组", "90g×8杯×12组", "杯", "＋3天", "全角加号"),
    ("6970618573019", "简爱酸奶吸吸乐黄皮百香果味100克x40袋（通路版）", "40", "袋", "+3到货", "旧写法兼容"),
    ("6923644288006", "现代牧场鲜牛奶百利包180ML*12袋", "180ML*12袋", "袋", "", "留空=不清零(原值3)"),
    ("6941704443776", "新希望(南山)塑瓶活润轻食瓶液体沙拉羽衣甘蓝牛油果风味发酵乳340G", "340G", "瓶", "随便写", "非法值=不写(原值0)"),
    ("6923061002988", "*友芝友金嚼鲜酪原风味发酵乳预制杯10g*5杯*8条", "10g*5杯*8条", "杯", "+3天", "停用商品→off-archive行"),
]

shutil.copy2(TPL, OUT)
wb = load_workbook(OUT)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
hdr = {str(h).strip() if h else "": i + 1 for i, h in enumerate(headers)}
print("模版列头:", headers)

col_bc = hdr["商品条码*"]
col_nm = next(v for k, v in hdr.items() if k.startswith("商品名称"))
col_sp = next(v for k, v in hdr.items() if k.startswith("规格"))
col_un = next(v for k, v in hdr.items() if k.startswith("单位"))
col_cy = next(v for k, v in hdr.items() if k.startswith("到货周期"))
col_cust = hdr.get("美联保康")
print("列号: 条码=%s 名称=%s 规格=%s 单位=%s 到货周期=%s 客户列=%s" % (col_bc, col_nm, col_sp, col_un, col_cy, col_cust))

# 清掉模版的示例行，写我们的数据行
ws.delete_rows(2, ws.max_row)
for i, (bc, nm, sp, un, cy, note) in enumerate(ROWS):
    r = 2 + i
    ws.cell(r, col_bc).value = bc
    ws.cell(r, col_nm).value = nm
    ws.cell(r, col_sp).value = sp
    ws.cell(r, col_un).value = un
    ws.cell(r, col_cy).value = cy
    ws.cell(r, col_cust).value = 5
wb.save(OUT)
print("已写 %d 行 → %s" % (len(ROWS), OUT))


def post_multipart(url, fields, files, headers):
    boundary = "----v184probe"
    body = b""
    for k, v in fields.items():
        body += ("--%s\r\nContent-Disposition: form-data; name=\"%s\"\r\n\r\n%s\r\n" % (boundary, k, v)).encode()
    for k, (fn, data) in files.items():
        body += ("--%s\r\nContent-Disposition: form-data; name=\"%s\"; filename=\"%s\"\r\n"
                 "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
                 % (boundary, k, fn)).encode()
        body += data + b"\r\n"
    body += ("--%s--\r\n" % boundary).encode()
    req = urllib.request.Request(url, data=body, headers={**headers,
        "Content-Type": "multipart/form-data; boundary=" + boundary})
    with urllib.request.urlopen(req, timeout=120) as resp:
        return resp.status, resp.read().decode("utf-8", "replace")


H = {"Authorization": "Bearer " + TOKEN, "X-Tenant-Id": TENANT}
with open(OUT, "rb") as f:
    raw = f.read()
st, txt = post_multipart(BASE + "/api/import/preview", {"category": "forecast_cross"},
                         {"file": (os.path.basename(OUT), raw)}, H)
print("\n=== preview HTTP", st, "===")
with open("/tmp/preview_9999.json", "w") as f:
    f.write(txt)
d = json.loads(txt)
print("顶层键:", list(d.keys()))
print("success =", d.get("success"), " detail =", d.get("detail"))
print("\nsuggestions:")
for s in (d.get("suggestions") or []):
    print("   ", json.dumps(s, ensure_ascii=False))
print("\ncross:")
print("   ", json.dumps(d.get("cross"), ensure_ascii=False)[:1200])
