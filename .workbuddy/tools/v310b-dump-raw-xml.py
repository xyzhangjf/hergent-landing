# -*- coding: utf-8 -*-
"""v310b 侦察：绕过 openpyxl 的 dimension 声明，直接读 sheet XML
"""
import os
import re
import zipfile

D = "/Users/zhangjunfeng/Documents/流水对账/舟谱导出的单据"
FILES = [
    "20260801-20260831采购单明细.xlsx",
    "2026年8月收入明细表.xlsx",
    "2026年8月费用明细表.xlsx",
]


def main():
    for name in FILES:
        p = os.path.join(D, name)
        print("=" * 78)
        print("文件:", name, "(%.1f KB)" % (os.path.getsize(p) / 1024))
        print("=" * 78)
        z = zipfile.ZipFile(p)
        # 1) dimension 声明
        x = z.read("xl/worksheets/sheet1.xml").decode("utf-8", "ignore")
        m = re.search(r"<dimension[^>]*/>", x)
        print("  <dimension> 声明:", m.group(0) if m else "(无)")
        # 2) 真实行列规模
        rows = re.findall(r'<row[^>]*r="(\d+)"', x)
        print("  <row> 实际条数:", len(rows),
              " 首/末:", (rows[0], rows[-1]) if rows else "-")
        cells = re.findall(r'<c[^>]*r="([A-Z]+)(\d+)"', x)
        cols = sorted({c[0] for c in cells}, key=lambda s: (len(s), s))
        print("  出现过的列:", cols)
        # 3) sharedStrings 规模
        ss = []
        if "xl/sharedStrings.xml" in z.namelist():
            s = z.read("xl/sharedStrings.xml").decode("utf-8", "ignore")
            ss = re.findall(r"<si>(.*?)</si>", s, re.S)
        print("  sharedStrings 条数:", len(ss))
        # 4) 用非 read_only 方式真正读取
        import openpyxl
        wb = openpyxl.load_workbook(p, data_only=True)
        ws = wb.worksheets[0]
        print("  正常模式 读到: max_row=%s max_col=%s" % (ws.max_row, ws.max_column))
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=6,
                                             max_col=min(ws.max_column or 0, 28))):
            vals = []
            for c in row:
                v = "" if c.value is None else str(c.value).replace("\n", "\\n").strip()
                vals.append(v)
            while vals and vals[-1] == "":
                vals.pop()
            print("    R%-2d | %s" % (i + 1, " | ".join(vals)))
        wb.close()
        print()


if __name__ == "__main__":
    main()
