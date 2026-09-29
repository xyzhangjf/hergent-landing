# -*- coding: utf-8 -*-
"""v313c M0.2 验收：放开「客户侧 + 内部侧」的钱单据导入。

## 与 M0.1 验收的差别（为什么不能再跑一遍旧的）
M0.1 只导供应商侧，落单时 `type` 是**写死** `supplier` 的。M0.2 要让
supplier / customer / internal 各就其位 ⇒ 关键是验**没有串池子**：

  · 客户费用单必须落到**客户**档案（不能拿客户名去供应商池找 ⇒ 那样会 100% 落空，
    看起来像「客户档案全缺」，实际是找错了池子）
  · 内部单往来单位是 `-`（本来就没有外部往来单位）⇒ 挂「内部往来」档案，不能写 0
  · `expense_orders.type` / `income_orders.type` 列必须逐单正确

## 三重隔离（与 v313-archive-gap.py 同一套）
1. 落库用**内存库** ⇒ 物理上不可能写真实库；
2. `ERP_DB_PATH` 在 import 前指向 /tmp 影子目录 ⇒ 启动期 schema 同步不碰 dev 库；
3. 期望值**从 xlsx 独立数出**，不照抄代码读数（照抄 = 探针跟着 bug 一起错）。

用法：
    cd /Users/zhangjunfeng/Documents/hergent-erp/server
    /usr/bin/python3 <本脚本>
"""
import collections
import json
import os
import shutil
import sqlite3
import sys

SRV = "/Users/zhangjunfeng/Documents/hergent-erp/server"
SNAP = "/tmp/v313-tenant1-snapshot.json"
SRC = "/Users/zhangjunfeng/Documents/流水对账/舟谱导出的单据"
FILES = {
    "zhoupu_income": os.path.join(SRC, "2026年8月收入明细表.xlsx"),
    "zhoupu_expense": os.path.join(SRC, "2026年8月费用明细表.xlsx"),
}
INTERNAL_CONTACT = "内部往来"

# ---------- 隔离（必须在 import 之前）----------
SHADOW_DIR = "/tmp/v313c-m02-shadow"
if os.path.isdir(SHADOW_DIR):
    shutil.rmtree(SHADOW_DIR)
os.makedirs(SHADOW_DIR)
shutil.copy2(os.path.join(SRV, "erp.db"), os.path.join(SHADOW_DIR, "erp.db"))
os.environ["ERP_DB_PATH"] = os.path.join(SHADOW_DIR, "erp.db")
os.environ.setdefault("ERP_SECRET", "v313c-m02-not-a-real-key")
sys.path.insert(0, SRV)

import routers.zhoupu_documents as Z  # noqa: E402

OK, BAD = [], []


def check(label, cond, extra=""):
    (OK if cond else BAD).append(label)
    print("   %s %s%s" % ("✅" if cond else "❌", label, ("  " + extra) if extra else ""))
    return cond


# ---------- 期望值：**从 xlsx 独立数出** ----------
def read_xlsx(fp):
    import openpyxl
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    ws.reset_dimensions()          # 舟谱 dimension 是假的，不 reset 会静默返回 0 行
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def expect(fp):
    """（单数, 行数, 金额）按单据类型分组 —— 直接读文件，不用代码读数当基准。"""
    rows = read_xlsx(fp)
    H = [str(x or "").strip() for x in rows[4]]
    i_no, i_ty = H.index("单据号"), H.index("单据类型")
    i_am = H.index([h for h in H if h.startswith("单据金额")][0])
    g = collections.OrderedDict()
    for r in rows[5:]:
        if not r or not r[i_no] or not str(r[i_no]).strip():
            continue
        no = str(r[i_no]).strip()
        if len(no) < 6:            # 「合计」尾行
            continue
        ty = str(r[i_ty] or "").strip()
        e = g.setdefault(ty, {"orders": set(), "lines": 0, "amount": 0.0})
        e["orders"].add(no)
        e["lines"] += 1
        try:
            e["amount"] += float(r[i_am] or 0)
        except Exception:
            pass
    return g


# ---------- 内存库 ----------
DDL_EXTRA = """
CREATE TABLE IF NOT EXISTS income_orders (
  id INTEGER PRIMARY KEY AUTOINCREMENT, order_no TEXT NOT NULL UNIQUE,
  type TEXT NOT NULL DEFAULT 'customer', contact_id INTEGER NOT NULL,
  category_id INTEGER DEFAULT 0, amount REAL DEFAULT 0, received_amount REAL DEFAULT 0,
  status TEXT DEFAULT 'unpaid', order_date TEXT DEFAULT '', note TEXT DEFAULT '',
  operator_id TEXT DEFAULT '', created_at TEXT DEFAULT '');
CREATE TABLE IF NOT EXISTS expense_orders (
  id INTEGER PRIMARY KEY AUTOINCREMENT, order_no TEXT NOT NULL UNIQUE,
  type TEXT NOT NULL DEFAULT 'customer', contact_id INTEGER NOT NULL,
  category_id INTEGER DEFAULT 0, amount REAL DEFAULT 0, paid_amount REAL DEFAULT 0,
  status TEXT DEFAULT 'unpaid', order_date TEXT DEFAULT '', note TEXT DEFAULT '',
  operator_id TEXT DEFAULT '', brand TEXT DEFAULT '', employee_id TEXT DEFAULT '',
  created_at TEXT DEFAULT '');
CREATE TABLE IF NOT EXISTS receivables (
  id INTEGER PRIMARY KEY AUTOINCREMENT, type TEXT NOT NULL DEFAULT 'ar',
  contact_id INTEGER NOT NULL, ref_type TEXT DEFAULT '', ref_id INTEGER DEFAULT 0,
  amount REAL DEFAULT 0, paid_amount REAL DEFAULT 0, status TEXT DEFAULT 'unpaid',
  due_date TEXT DEFAULT '', created_at TEXT DEFAULT '');
"""


def mem_db(with_internal=True):
    snap = json.load(open(SNAP, encoding="utf-8"))
    c = sqlite3.connect(":memory:")
    c.row_factory = sqlite3.Row
    for t, v in snap["_tables"].items():
        cols = v["cols"]
        c.execute('CREATE TABLE %s (%s)' % (t, ", ".join('"%s"' % x for x in cols)))
        c.executemany("INSERT INTO %s VALUES (%s)" % (t, ", ".join("?" * len(cols))),
                      v["rows"])
    c.executescript(DDL_EXTRA)
    if with_internal:
        cols = snap["_tables"]["contacts"]["cols"]
        row = {k: "" for k in cols}
        row["id"] = 999001
        row["name"] = INTERNAL_CONTACT
        row["type"] = "internal"
        row["is_active"] = 1
        c.execute("INSERT INTO contacts VALUES (%s)" % ", ".join("?" * len(cols)),
                  [row.get(k) for k in cols])
    c.commit()
    return c


def run(kind, fp, with_internal=True):
    conn = mem_db(with_internal)
    if not with_internal:
        # 🔴 撤掉「内部往来」必须**真的删掉这一行**：快照里本来就有它（生产已建，id=2973），
        #    光靠「不插入模拟行」是撤不掉的 —— 那会让 T4 恒绿（假绿）。
        conn.execute("DELETE FROM contacts WHERE name=?", (INTERNAL_CONTACT,))
        conn.commit()
    # 走 `run_file`：它自己判别形态 → 建列绑定 → 按 family 分派落库路径。
    #   不手拼 detect/bind ⇒ 与用户在界面上点「导入」走的是**同一条代码路径**。
    rep = Z.run_file(conn, fp, dry_run=False)
    st = rep["stats"]
    # 落库后按 type 分组核对
    tbl = "income_orders" if kind == "zhoupu_income" else "expense_orders"
    got = {}
    for r in conn.execute("SELECT type, COUNT(*) n, ROUND(SUM(amount),2) amt "
                          "FROM %s GROUP BY type" % tbl):
        got[r["type"]] = (r["n"], r["amt"])
    conn.close()
    return st, got


def main():
    print("=" * 78)
    print("v313c M0.2 验收：客户侧 + 内部侧放开")
    print("=" * 78)

    # ---------- T1 期望值来自文件 ----------
    print()
    print("【T1】期望值（从 xlsx 独立数出，不照抄代码读数）")
    EXP = {}
    for kind, fp in FILES.items():
        EXP[kind] = expect(fp)
        print("   %s：" % kind)
        for ty, e in EXP[kind].items():
            print("      %-8s %3d 单 / %4d 行 / ¥%s"
                  % (ty, len(e["orders"]), e["lines"], format(e["amount"], ",.2f")))

    TYPE_OF = {"供应商收入单": "supplier", "客户收入单": "customer", "内部收入单": "internal",
               "供应商费用单": "supplier", "客户费用单": "customer", "内部费用单": "internal"}

    # ---------- T2 真落库（含「内部往来」档案）----------
    print()
    print("【T2】真落库 —— 三类各就其位")
    for kind, fp in FILES.items():
        st, got = run(kind, fp, with_internal=True)
        tbl = "income" if kind == "zhoupu_income" else "expense"
        e_tot = sum(len(v["orders"]) for v in EXP[kind].values())
        print("   %s: total=%d created=%d blocked=%d failed=%d existed=%d"
              % (kind, st["orders_total"], st["orders_created"], st["orders_blocked"],
                 st["orders_failed"], st["orders_existed"]))
        check("T2 %s 张数齐（期望 %d）" % (kind, e_tot), st["orders_created"] == e_tot,
              "实际 %d" % st["orders_created"])
        check("T2 %s 零阻塞零失败" % kind,
              st["orders_blocked"] == 0 and st["orders_failed"] == 0,
              "blocked=%d failed=%d" % (st["orders_blocked"], st["orders_failed"]))
        # 逐类型核对张数与金额
        for ty, e in EXP[kind].items():
            bt = TYPE_OF[ty]
            n, amt = got.get(bt, (0, 0.0))
            check("T2 %s 落库 type=%s：%d 单" % (tbl, bt, len(e["orders"])),
                  n == len(e["orders"]), "实际 %d 单" % n)
            check("T2 %s 落库 type=%s：金额 ¥%s" % (tbl, bt, format(e["amount"], ",.2f")),
                  abs((amt or 0.0) - e["amount"]) < 0.01,
                  "实际 ¥%s" % format(amt or 0.0, ",.2f"))

    # ---------- T3 没有串池子（反例）----------
    print()
    print("【T3】反例：客户单不能落到供应商档案上")
    conn = mem_db(True)
    cust_ids = set(r["id"] for r in conn.execute(
        "SELECT id FROM contacts WHERE type='customer'"))
    sup_ids = set(r["id"] for r in conn.execute(
        "SELECT id FROM contacts WHERE type='supplier'"))
    Z.run_file(conn, FILES["zhoupu_expense"], dry_run=False)
    rows = list(conn.execute(
        "SELECT type, contact_id FROM expense_orders"))
    bad_cust_on_sup = [r for r in rows
                       if r["type"] == "customer" and r["contact_id"] in sup_ids]
    bad_cust_not_cust = [r for r in rows
                         if r["type"] == "customer" and r["contact_id"] not in cust_ids]
    check("T3 无「客户单挂在供应商档案」", not bad_cust_on_sup,
          "%d 张" % len(bad_cust_on_sup))
    check("T3 客户单 contact_id 全部落在客户档案", not bad_cust_not_cust,
          "越界 %d 张" % len(bad_cust_not_cust))
    # 「内部往来」的 id **从库里现查**（不能写死：本地模拟是 999001、生产真实是 2973）
    ic_id = conn.execute("SELECT id FROM contacts WHERE name=?",
                         (INTERNAL_CONTACT,)).fetchone()["id"]
    internal = [r for r in rows if r["type"] == "internal"]
    check("T3 内部单全部挂到「%s」档案（id=%d）" % (INTERNAL_CONTACT, ic_id),
          bool(internal) and all(r["contact_id"] == ic_id for r in internal),
          "%d 张，contact_id 集合=%s" % (len(internal),
                                     sorted(set(r["contact_id"] for r in internal))[:5]))
    check("T3 没有 contact_id=0（孤儿行）",
          not [r for r in rows if not r["contact_id"]],
          "0 值 %d 张" % len([r for r in rows if not r["contact_id"]]))
    conn.close()

    # ---------- T4 判别力自证：撤掉「内部往来」档案 ----------
    print()
    print("【T4】判别力自证：撤掉「%s」档案后内部单必须**可见地**被挡" % INTERNAL_CONTACT)
    st4, got4 = run("zhoupu_expense", FILES["zhoupu_expense"], with_internal=False)
    n_internal = len(EXP["zhoupu_expense"].get("内部费用单", {"orders": set()})["orders"])
    check("T4 内部单被挡（期望 blocked=%d）" % n_internal,
          st4["orders_blocked"] == n_internal, "实际 blocked=%d" % st4["orders_blocked"])
    check("T4 被挡的张数**仍计入总数**（不许从报告里消失）",
          st4["orders_total"] == sum(len(v["orders"]) for v in EXP["zhoupu_expense"].values()),
          "total=%d" % st4["orders_total"])
    check("T4 客户 / 供应商侧不受影响（照旧落库）",
          got4.get("customer", (0, 0))[0] == len(
              EXP["zhoupu_expense"].get("客户费用单", {"orders": set()})["orders"]),
          "customer=%d" % got4.get("customer", (0, 0))[0])

    print()
    print("=" * 78)
    print("通过 %d / 失败 %d" % (len(OK), len(BAD)))
    if BAD:
        for b in BAD:
            print("   ❌ " + b)
        sys.exit(1)
    print("全部通过")


if __name__ == "__main__":
    main()
