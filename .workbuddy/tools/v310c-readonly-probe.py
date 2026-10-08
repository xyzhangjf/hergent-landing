# -*- coding: utf-8 -*-
"""v310c 判别力自证：read_only=True 在「dimension 撒谎」的表上会读到几行？

正例：read_only=False（不信任 dimension，直接扫 XML）→ 应得真实行数
反例：read_only=True （信任 dimension）→ 猜会只读 1 行
若两者不一致 ⇒ 现有导入管线（全部 read_only=True）读这类文件是**静默 0 行**。
"""
import io
import os
import warnings
import openpyxl

warnings.filterwarnings("ignore")

D = "/Users/zhangjunfeng/Documents/流水对账/舟谱导出的单据"
FILES = [
    ("20260801-20260831采购单明细.xlsx", 1140),
    ("2026年8月收入明细表.xlsx", 20),
    ("2026年8月费用明细表.xlsx", 240),
]

print("%-34s %-12s %-12s %-10s %s" % ("文件", "期望行数", "ro=True", "ro=False", "判定"))
print("-" * 92)
bad = 0
for name, expect in FILES:
    p = os.path.join(D, name)
    raw = open(p, "rb").read()
    res = {}
    for mode in (True, False):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=mode, data_only=True)
            ws = wb.worksheets[0]
            n = ws.max_row or 0
            # 真读一遍统计非空行，避免 max_row 也撒谎
            cnt = 0
            for r in ws.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() != "" for v in r):
                    cnt += 1
            res[mode] = (n, cnt)
            wb.close()
        except Exception as e:
            res[mode] = ("ERR:%s" % type(e).__name__, -1)
    rt, rf = res[True], res[False]
    ok = (rf[1] == expect)
    verdict = "OK 正例达标" if ok else "!! 期望 %d 实得 %d" % (expect, rf[1])
    if rt[1] != rf[1]:
        verdict += "  ⇒ read_only 丢 %d 行" % (rf[1] - rt[1])
        bad += 1
    print("%-34s %-12s %-12s %-10s %s" % (
        name[:32], expect,
        "max=%s/实=%s" % rt, "max=%s/实=%s" % rf, verdict))
print("-" * 92)
print("结论：%s" % ("read_only=True 在被测文件上【确实丢行】⇒ 现有管线静默失效"
                 if bad else "read_only=True 未丢行（本组样本）"))
